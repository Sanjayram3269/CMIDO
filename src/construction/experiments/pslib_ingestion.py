from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl

from .dataset import build_canonical_dataset


PSLIB_INGESTION_SCHEMA_VERSION = "9B-1.0"
SUPPORTED_PSLIB_VERSION = "3.4"
_DURATION_RE = re.compile(r"^\s*(?P<value>\d+(?:\.\d+)?)\s*(?P<unit>[a-zA-Z]+)\s*$")
_DEP_RE = re.compile(r"^\s*(?P<id>\d+)\s*(?P<rel>FS|SS|FF|SF)\s*$", re.IGNORECASE)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _duration_days(raw: Any) -> int:
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        value = float(raw)
        unit = "day"
    else:
        match = _DURATION_RE.match(str(raw))
        if not match:
            raise ValueError(f"unsupported PSLIB duration: {raw!r}")
        value = float(match.group("value"))
        unit = match.group("unit").lower()
    factors = {"d": 1.0, "day": 1.0, "days": 1.0, "h": 1 / 24, "hour": 1 / 24, "hours": 1 / 24, "w": 7.0, "week": 7.0, "weeks": 7.0}
    if unit not in factors:
        raise ValueError(f"unsupported PSLIB duration unit: {unit}")
    converted = value * factors[unit]
    if converted < 0 or not converted.is_integer():
        raise ValueError(f"duration does not convert to integer CMIDO days: {raw!r}")
    return int(converted)


def _parse_predecessors(raw: Any) -> list[tuple[str, str]]:
    if raw in (None, ""):
        return []
    result: list[tuple[str, str]] = []
    for token in str(raw).split(";"):
        token = token.strip()
        if not token:
            continue
        match = _DEP_RE.match(token)
        if not match:
            raise ValueError(f"unsupported PSLIB predecessor token: {token!r}")
        result.append((match.group("id"), match.group("rel").upper()))
    return result


def _sheet_rows(ws, header_row: int) -> list[dict[str, Any]]:
    headers = [cell.value for cell in ws[header_row]]
    rows: list[dict[str, Any]] = []
    for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(value not in (None, "") for value in values):
            continue
        rows.append({str(headers[i]): _json_value(values[i]) for i in range(len(headers)) if headers[i] not in (None, "")})
    return rows


def ingest_pslib_project(path: str | Path, *, dataset_id: str | None = None) -> dict[str, Any]:
    """Convert one empirical PSLIB v3.4 Excel project into the CMIDO canonical project envelope."""
    source = Path(path).resolve()
    if not source.exists() or source.suffix.lower() != ".xlsx":
        raise ValueError(f"PSLIB project workbook does not exist or is not .xlsx: {source}")

    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    required_sheets = {"Baseline Schedule", "Resources", "Risk Analysis", "Agenda"}
    missing = required_sheets - set(workbook.sheetnames)
    if missing:
        raise ValueError(f"PSLIB workbook missing required sheets: {sorted(missing)}")

    schedule = workbook["Baseline Schedule"]
    schedule_rows = list(schedule.iter_rows(min_row=3, values_only=True))
    if not schedule_rows:
        raise ValueError("PSLIB Baseline Schedule contains no records")

    project_row = schedule_rows[0]
    project_id = source.stem
    project_name = str(project_row[1] or project_id)

    activities: list[dict[str, Any]] = []
    dependencies: list[dict[str, Any]] = []
    rejected_dependencies: list[dict[str, Any]] = []
    activity_source_rows: dict[str, int] = {}

    for excel_row, row in enumerate(schedule_rows, start=3):
        activity_id = row[0]
        name = row[1]
        duration_raw = row[7] if len(row) > 7 else None
        if not isinstance(activity_id, (int, float)) or isinstance(activity_id, bool) or not name:
            continue
        if duration_raw in (None, ""):
            continue
        activity_id = str(int(activity_id)) if float(activity_id).is_integer() else str(activity_id)
        try:
            duration_days = _duration_days(duration_raw)
        except ValueError as exc:
            raise ValueError(f"{source.name}: row {excel_row}: {exc}") from exc

        predecessors_raw = row[3] if len(row) > 3 else None
        baseline_start = row[5] if len(row) > 5 else None
        baseline_end = row[6] if len(row) > 6 else None
        resource_demand = row[8] if len(row) > 8 else None
        fixed_cost = row[10] if len(row) > 10 else None
        cost_per_hour = row[11] if len(row) > 11 else None
        variable_cost = row[12] if len(row) > 12 else None
        total_cost = row[13] if len(row) > 13 else None

        activities.append({
            "activity_id": activity_id,
            "project_id": project_id,
            "activity_code": activity_id,
            "activity_name": str(name),
            "duration_days": duration_days,
            "wbs": _json_value(row[2] if len(row) > 2 else None),
            "source_duration": _json_value(duration_raw),
            "planned_start": _json_value(baseline_start),
            "planned_finish": _json_value(baseline_end),
            "resource_demand_raw": _json_value(resource_demand),
            "fixed_cost": _json_value(fixed_cost),
            "cost_per_hour": _json_value(cost_per_hour),
            "variable_cost": _json_value(variable_cost),
            "total_cost": _json_value(total_cost),
            "source_row": excel_row,
            "provenance": "OBS",
        })
        activity_source_rows[activity_id] = excel_row

        try:
            parsed = _parse_predecessors(predecessors_raw)
        except ValueError as exc:
            rejected_dependencies.append({"activity_id": activity_id, "source_row": excel_row, "field": "Predecessors", "reason": str(exc)})
            continue
        for predecessor_id, relationship in parsed:
            dependencies.append({
                "dependency_id": f"D{len(dependencies) + 1:05d}",
                "project_id": project_id,
                "predecessor_id": predecessor_id,
                "successor_id": activity_id,
                "relationship_type": relationship,
                "lag_days": 0,
                "provenance": "OBS",
            })

    activity_ids = {item["activity_id"] for item in activities}
    missing_refs = [d for d in dependencies if d["predecessor_id"] not in activity_ids or d["successor_id"] not in activity_ids]
    if missing_refs:
        raise ValueError(
            f"{source.name}: {len(missing_refs)} observed predecessor references point to rows without task durations; "
            "the source structure must be reviewed before CMIDO can canonicalize it"
        )

    canonical = build_canonical_dataset(
        {"project_id": project_id, "project_name": project_name, "project_type": "empirical_pslib"},
        activities,
        dependencies,
        source=str(source),
        dataset_id=dataset_id or f"CMIDO_9B_{project_id}",
        metadata={
            "ingestion_schema_version": PSLIB_INGESTION_SCHEMA_VERSION,
            "source_format": "PSLIB_XLSX",
            "pslib_release": SUPPORTED_PSLIB_VERSION,
            "source_sha256": _sha256(source),
            "source_workbook_sheets": workbook.sheetnames,
            "activity_source_rows": activity_source_rows,
            "rejected_dependency_records": rejected_dependencies,
        },
    )

    risk_rows = _sheet_rows(workbook["Risk Analysis"], 2)
    risk_by_id: dict[str, dict[str, Any]] = {}
    for row in risk_rows:
        raw_id = row.get("ID")
        if isinstance(raw_id, (int, float)) and not isinstance(raw_id, bool):
            risk_by_id[str(int(raw_id)) if float(raw_id).is_integer() else str(raw_id)] = {
                "description": row.get("Description"),
                "optimistic_raw": row.get("Optimistic"),
                "most_probable_raw": row.get("Most Probable"),
                "pessimistic_raw": row.get("Pessimistic"),
                "provenance": "OBS",
            }
    for activity in canonical["activities"]:
        activity["risk_profile_observed"] = risk_by_id.get(activity["activity_id"])

    resources = _sheet_rows(workbook["Resources"], 2)
    calendar_rows = _sheet_rows(workbook["Agenda"], 1)
    control_snapshots: dict[str, Any] = {}
    for sheet_name in ("Project Control - TP1", "TP2", "TP3"):
        if sheet_name not in workbook.sheetnames:
            continue
        ws = workbook[sheet_name]
        status_date = _json_value(ws.cell(1, 3).value)
        rows = _sheet_rows(ws, 4)
        control_snapshots[sheet_name] = {"status_date": status_date, "activity_rows": rows, "provenance": "OBS"}

    canonical["metadata"].update({
        "resource_rows": resources,
        "calendar_rows": calendar_rows,
        "risk_profile_count": len(risk_by_id),
        "control_snapshot_names": sorted(control_snapshots),
        "control_snapshots": control_snapshots,
    })
    canonical["fingerprint_sha256"] = hashlib.sha256(
        json.dumps(canonical, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode("utf-8")
    ).hexdigest()
    return canonical


def build_pslib_audit(dataset: dict[str, Any]) -> dict[str, Any]:
    metadata = dataset.get("metadata", {})
    return {
        "schema_version": PSLIB_INGESTION_SCHEMA_VERSION,
        "dataset_id": dataset.get("dataset_id"),
        "project_id": dataset.get("project", {}).get("project_id"),
        "project_name": dataset.get("project", {}).get("project_name"),
        "pslib_release": metadata.get("pslib_release"),
        "source": dataset.get("source"),
        "source_sha256": metadata.get("source_sha256"),
        "canonical_fingerprint_sha256": dataset.get("fingerprint_sha256"),
        "activity_count": len(dataset.get("activities", [])),
        "dependency_count": len(dataset.get("dependencies", [])),
        "risk_profile_count": metadata.get("risk_profile_count"),
        "control_snapshot_names": metadata.get("control_snapshot_names", []),
        "resource_row_count": len(metadata.get("resource_rows", [])),
        "rejected_dependency_records": len(metadata.get("rejected_dependency_records", [])),
    }
